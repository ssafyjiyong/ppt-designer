import io
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

SLIDE_W = 13.333
SLIDE_H = 7.5

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_SHAPE = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
}
_CHART_TYPE = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def _hex_to_rgb(h: str) -> RGBColor:
    if not h:
        return RGBColor(0, 0, 0)
    h = h.lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", h):
        return RGBColor(0, 0, 0)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def _blank_layout(prs: Presentation):
    return prs.slide_layouts[6]


def _add_text(slide, el: dict):
    box = slide.shapes.add_textbox(
        Inches(float(el["x"])), Inches(float(el["y"])),
        Inches(float(el["w"])), Inches(float(el["h"])),
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = _ANCHOR.get(el.get("valign", "top"), MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    p.alignment = _ALIGN.get(el.get("align", "left"), PP_ALIGN.LEFT)

    base_size = int(el.get("font_size", 16))
    base_color = el.get("color", "#1A1A1A")
    base_bold = bool(el.get("bold", False))
    base_italic = bool(el.get("italic", False))
    font_name = el.get("font_name", "Pretendard")

    # runs: 한 줄 안에서 부분별로 색/굵기를 다르게 (키워드 강조)
    runs = el.get("runs")
    if not runs:
        runs = [{"text": el.get("text", "")}]
    for r in runs:
        run = p.add_run()
        run.text = str(r.get("text", ""))
        f = run.font
        f.size = Pt(int(r.get("size", base_size)))
        f.bold = bool(r.get("bold", base_bold))
        f.italic = bool(r.get("italic", base_italic))
        f.color.rgb = _hex_to_rgb(r.get("color", base_color))
        f.name = font_name


def _add_bullet(slide, el: dict):
    box = slide.shapes.add_textbox(
        Inches(float(el["x"])), Inches(float(el["y"])),
        Inches(float(el["w"])), Inches(float(el["h"])),
    )
    tf = box.text_frame
    tf.word_wrap = True
    items = el.get("items", [])
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = f"• {item}"
        run.font.size = Pt(int(el.get("font_size", 14)))
        run.font.color.rgb = _hex_to_rgb(el.get("color", "#1A1A1A"))
        run.font.name = el.get("font_name", "Pretendard")


def _add_shape(slide, el: dict):
    shape_type = el.get("shape", "rectangle")
    if shape_type == "line":
        from pptx.util import Emu
        x1 = Inches(float(el["x"]))
        y1 = Inches(float(el["y"]))
        x2 = Inches(float(el["x"]) + float(el["w"]))
        y2 = Inches(float(el["y"]) + float(el["h"]))
        connector = slide.shapes.add_connector(1, x1, y1, x2, y2)
        connector.line.color.rgb = _hex_to_rgb(el.get("line") or "#1A1A1A")
        connector.line.width = Pt(float(el.get("line_w", 1)))
        return
    mso_shape = _SHAPE.get(shape_type, MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(
        mso_shape,
        Inches(float(el["x"])), Inches(float(el["y"])),
        Inches(float(el["w"])), Inches(float(el["h"])),
    )
    fill = el.get("fill")
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _hex_to_rgb(fill)
    else:
        shp.fill.background()
    line = el.get("line")
    if line:
        shp.line.color.rgb = _hex_to_rgb(line)
        shp.line.width = Pt(float(el.get("line_w", 1)))
        if el.get("dash"):
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False


def _add_table(slide, el: dict):
    rows = el.get("rows") or []
    if not rows:
        return
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows)
    gshape = slide.shapes.add_table(
        n_rows, n_cols,
        Inches(float(el["x"])), Inches(float(el["y"])),
        Inches(float(el["w"])), Inches(float(el["h"])),
    )
    table = gshape.table
    has_header = bool(el.get("header", True))
    table.first_row = has_header
    size = int(el.get("font_size", 14))
    header_fill = el.get("header_fill", "#1970DF")
    header_color = el.get("header_color", "#FFFFFF")
    body_color = el.get("body_color", "#172A3D")
    body_fill = el.get("body_fill")
    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            cell = table.cell(ri, ci)
            cell.text = str(row[ci]) if ci < len(row) else ""
            is_head = has_header and ri == 0
            if is_head:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _hex_to_rgb(header_fill)
            elif body_fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _hex_to_rgb(body_fill)
            else:
                cell.fill.background()
            para = cell.text_frame.paragraphs[0]
            run = para.runs[0] if para.runs else para.add_run()
            f = run.font
            f.size = Pt(size)
            f.bold = is_head
            f.name = el.get("font_name", "Pretendard")
            f.color.rgb = _hex_to_rgb(header_color if is_head else body_color)


def _add_chart(slide, el: dict):
    chart_type = _CHART_TYPE.get(el.get("chart_type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
    data = CategoryChartData()
    data.categories = el.get("categories") or []
    series = el.get("series") or []
    for s in series:
        data.add_series(s.get("name", ""), tuple(float(v) for v in s.get("values", [])))
    gframe = slide.shapes.add_chart(
        chart_type,
        Inches(float(el["x"])), Inches(float(el["y"])),
        Inches(float(el["w"])), Inches(float(el["h"])),
        data,
    )
    chart = gframe.chart
    chart.has_title = False
    show_legend = bool(el.get("legend", len(series) > 1))
    chart.has_legend = show_legend
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    colors = el.get("colors") or []
    if colors and chart_type == XL_CHART_TYPE.PIE and chart.plots:
        points = chart.plots[0].series[0].points
        for i, pt in enumerate(points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _hex_to_rgb(colors[i % len(colors)])
    elif colors:
        for i, plot_series in enumerate(chart.series):
            plot_series.format.fill.solid()
            plot_series.format.fill.fore_color.rgb = _hex_to_rgb(colors[i % len(colors)])


def _apply_background(slide, bg: dict):
    color_hex = (bg or {}).get("color", "#FFFFFF")
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _hex_to_rgb(color_hex)


def _render_slide(slide, spec: dict):
    _apply_background(slide, spec.get("background", {}))
    for el in spec.get("elements", []):
        try:
            t = el.get("type")
            if t == "text":
                _add_text(slide, el)
            elif t == "bullet":
                _add_bullet(slide, el)
            elif t == "shape":
                _add_shape(slide, el)
            elif t == "table":
                _add_table(slide, el)
            elif t == "chart":
                _add_chart(slide, el)
        except Exception as e:
            print(f"[render] element skip: {e} / el={el}")


def build_pptx(specs: list) -> bytes:
    prs = _new_presentation()
    layout = _blank_layout(prs)
    for spec in specs:
        slide = prs.slides.add_slide(layout)
        _render_slide(slide, spec or {})
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def extract_drafts(pptx_bytes: bytes) -> list:
    prs = Presentation(io.BytesIO(pptx_bytes))
    drafts = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                line = "".join(r.text for r in p.runs).strip()
                if line:
                    parts.append(line)
        drafts.append("\n".join(parts))
    return drafts
